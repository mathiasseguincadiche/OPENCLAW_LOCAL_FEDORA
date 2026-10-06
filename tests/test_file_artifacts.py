from __future__ import annotations

import io
import json
import re
import threading
from http.client import HTTPConnection
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit

import pytest
from docx import Document
from pypdf import PdfReader
from test_daily_worker import planned as planned
from test_guided_learning import reviewer

from clawfedora.agent_tools import invoke
from clawfedora.agents import deploy_workspaces
from clawfedora.document_exports import EXPORT_PREFIX, export, output_bytes
from clawfedora.file_artifacts import produce, read_artifact
from clawfedora.learning import checkpoints, initialize, submit
from clawfedora.project_common import read_json, write_json
from clawfedora.project_worker import _collect, _resolve_tool_files, run_project_tasks
from clawfedora.webui_bridge import make_server

ROOT = Path(__file__).resolve().parents[1]
SOURCE = (
    "# Déploiement OPS\n\nVérifier la stratégie et le retour arrière.\n\n"
    "| Outil | Rôle |\n|---|---|\n| Git | Versionner |\n\n```bash\nprintf 'test'\n```\n"
)


def test_actual_document_formats_share_content_and_are_reproducible() -> None:
    for kind in ("pdf", "docx", "txt"):
        raw = export(SOURCE, kind)
        assert export(SOURCE, kind) == raw
        if kind == "pdf":
            document = PdfReader(io.BytesIO(raw))
            assert len(document.pages) == 1
            assert "Déploiement OPS" in document.pages[0].extract_text()
            assert "Versionner" in document.pages[0].extract_text()
        elif kind == "docx":
            docx = Document(io.BytesIO(raw))
            assert docx.paragraphs[0].style.name == "Heading 1"
            assert docx.tables[0].cell(1, 1).text == "Versionner"
        else:
            assert "# Déploiement" not in raw.decode()
            assert "Git | Versionner" in raw.decode()


def test_untrusted_markdown_is_literal_without_external_reads(monkeypatch: Any) -> None:
    import socket

    def forbidden(*_args: Any, **_kwargs: Any) -> Any:
        raise AssertionError("network fetch forbidden")

    monkeypatch.setattr(socket, "create_connection", forbidden)
    source = '# Guide\n\n![secret](file:///etc/passwd)\n\n<img src="http://127.0.0.1/secret">\n'
    for kind in ("pdf", "docx", "txt"):
        assert export(source, kind)
    with pytest.raises(ValueError):
        export("a" * 60001, "pdf")


def test_artifact_references_bind_role_task_source_and_real_bytes(tmp_path: Path) -> None:
    deploy_workspaces(ROOT, tmp_path)
    role = "redacteur-pedagogique"
    workspace = tmp_path / "workspaces" / role
    result = invoke(
        tmp_path,
        role,
        workspace,
        "clawfedora_artifact",
        {
            "format": "markdown",
            "filename": "fiche-ops.md",
            "content": SOURCE,
            "exports": ["pdf", "docx", "txt"],
        },
    )
    response = {
        "files": {f"deliverables/guide/guide.{ext}": ref for ext, ref in result["references"].items()}
    }
    _resolve_tool_files(workspace, set(), {"role": role}, response)
    project = tmp_path / "project"
    project.mkdir()
    task = {"id": "guide", "role": role, "expected_outputs": list(response["files"])}
    _collect(ROOT, project, task, response)
    assert PdfReader(project / "deliverables/guide/guide.pdf").pages[0].extract_text()
    assert Document(project / "deliverables/guide/guide.docx").tables
    assert (project / "deliverables/guide/guide.md").read_text() == SOURCE
    reference = result["references"]["pdf"]
    for before, role_name, files in [
        ({workspace / result["receipt"]}, role, {"deliverables/guide/guide.pdf": reference}),
        (set(), "chef-operations", {"deliverables/guide/guide.pdf": reference}),
        (set(), role, {"deliverables/guide/guide.pdf": reference}),
    ]:
        with pytest.raises(ValueError):
            _resolve_tool_files(workspace, before, {"role": role_name}, {"files": files})
    identifier = reference.split(":")[1]
    for role_name, requested_id, extension in (
        (role, "../" + identifier, "pdf"),
        ("../redacteur-pedagogique", identifier, "pdf"),
        (role, identifier, "../../etc/passwd"),
    ):
        with pytest.raises(ValueError):
            read_artifact(workspace, role_name, requested_id, extension)
    path, _ = read_artifact(workspace, role, identifier, "pdf")
    path.chmod(0o640)
    path.write_bytes(b"fake")
    with pytest.raises(ValueError, match="modifié"):
        read_artifact(workspace, role, identifier, "pdf")


def test_code_templates_remain_literal_and_invalid_formats_are_refused(tmp_path: Path) -> None:
    deploy_workspaces(ROOT, tmp_path)
    devops = tmp_path / "workspaces/ingenieur-devops"
    for kind, content in [
        ("json", '{"example":true}\n'),
        ("template", "server={{ hostname }}\n"),
        ("hcl", 'resource "null_resource" "demo" {}\n'),
        ("dockerfile", "FROM scratch\n"),
        ("yaml", "---\na: 1\n---\nb: 2\n"),
        ("shell", "#!/bin/bash\ntouch NEVER_EXECUTED\n"),
    ]:
        before = set((devops / ".clawfedora-tool-evidence").glob("*.json"))
        result = produce(devops, "ingenieur-devops", {"format": kind, "content": content})
        ext, reference = next(iter(result["references"].items()))
        path, _ = read_artifact(devops, "ingenieur-devops", reference.split(":")[1], ext)
        assert path.read_text() == content
        from clawfedora.project_worker import _collect_tool_receipts
        destination = tmp_path / "project"
        destination.mkdir(exist_ok=True)
        _collect_tool_receipts(devops, before, destination, "configuration")
    assert not (devops / "NEVER_EXECUTED").exists()
    for data in (
        {"format": "json", "content": "invalid"},
        {
            "format": "xml",
            "content": '<!DOCTYPE a [<!ENTITY s SYSTEM "file:///etc/passwd">]><a>&s;</a>',
        },
        {"format": "yaml", "content": "a: 1", "exports": ["pdf"]},
    ):
        with pytest.raises(ValueError):
            produce(devops, "ingenieur-devops", data)
    with pytest.raises(ValueError, match="interdit"):
        produce(
            tmp_path / "workspaces/auditeur-qualite",
            "auditeur-qualite",
            {"format": "shell", "content": "true"},
        )
    with pytest.raises(ValueError):
        output_bytes("x.pdf", {"x.pdf": "fake PDF"})
    with pytest.raises(ValueError):
        output_bytes("x.pdf", {"x.pdf": EXPORT_PREFIX + "/etc/passwd"})


def test_guided_feedback_publishes_exact_reviewed_binary_exports(
    planned: tuple[Path, Path],
) -> None:
    runtime, project = planned
    task_path = project / "context/tasks/design-choice.json"
    packet = read_json(task_path)
    packet["task"]["expected_outputs"] = [
        f"deliverables/design-choice/report.{ext}" for ext in ("md", "pdf", "docx")
    ]
    write_json(task_path, packet)
    initialize(project, "guided")

    def starter(role: str, prompt: str, _session: str) -> dict[str, Any]:
        task = json.loads(prompt.split("\n", 1)[1])
        result = invoke(
            runtime,
            role,
            runtime / "workspaces" / role,
            "clawfedora_artifact",
            {"format": "markdown", "content": SOURCE, "exports": ["pdf", "docx"]},
        )
        return {
            "files": {p: result["references"][Path(p).suffix[1:]] for p in task["expected_outputs"]},
            "summary": "Compléter les preuves.",
        }

    run_project_tasks(ROOT, runtime, project, runner=starter)
    item = checkpoints(project)[0]
    assert not (project / "deliverables/design-choice/report.pdf").exists()
    files = dict(item["files"])
    files["deliverables/design-choice/report.md"] += "\nLimite : aucun déploiement exécuté.\n"
    submit(
        ROOT,
        runtime,
        project,
        item["task_id"],
        {"files": files, "explanation": "J’ai précisé la limite", "human_approved": True},
    )
    run_project_tasks(ROOT, runtime, project, runner=reviewer)
    assert checkpoints(project)[0]["status"] == "REVIEWED"
    pdf = project / "deliverables/design-choice/report.pdf"
    assert "aucun déploiement exécuté" in PdfReader(pdf).pages[0].extract_text()
    assert pdf.read_bytes() == output_bytes("deliverables/design-choice/report.pdf", files)


def test_chat_returns_signed_downloads_without_operator_token(tmp_path: Path) -> None:
    deploy_workspaces(ROOT, tmp_path)
    token = "a" * 64

    def runner(role: str, _prompt: str, _session: str) -> dict[str, Any]:
        invoke(
            tmp_path,
            role,
            tmp_path / "workspaces" / role,
            "clawfedora_artifact",
            {
                "format": "markdown",
                "filename": "fiche-ops.md",
                "content": SOURCE,
                "exports": ["pdf", "docx", "txt"],
            },
        )
        return {"text": "Voici le guide."}

    with make_server(ROOT, tmp_path, token, 0, runner=runner) as server:
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        conn = HTTPConnection("127.0.0.1", server.server_port)
        conn.request(
            "POST",
            "/v1/chat/completions",
            json.dumps(
                {
                    "model": "openclaw/redacteur-pedagogique",
                    "messages": [{"role": "user", "content": "Un guide PDF"}],
                }
            ),
            {"Authorization": "Bearer " + token, "Content-Type": "application/json"},
        )
        response = conn.getresponse()
        assert response.status == 200
        text = json.loads(response.read())["choices"][0]["message"]["content"]
        assert token not in text
        urls = re.findall(r"\]\((http://[^)]+)\)", text)
        assert len(urls) == 4
        pdf_url = next(urlsplit(url) for url in urls if "/pdf?" in url)
        conn.request("GET", pdf_url.path + "?" + pdf_url.query)
        result = conn.getresponse()
        assert result.status == 200 and result.read().startswith(b"%PDF-")
        assert "fiche-ops.pdf" in result.getheader("Content-Disposition")
        conn.request(
            "GET", pdf_url.path + "?" + pdf_url.query.replace("signature=", "signature=wrong")
        )
        result = conn.getresponse()
        assert result.status == 403
        result.read()
        conn.close()
        server.shutdown()
        thread.join(timeout=5)
