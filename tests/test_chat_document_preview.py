from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from reportlab.pdfgen import canvas

from clawfedora import chat_files
from clawfedora.chat_document_preview import preview_document


def _upload(runtime: Path, name: str, raw: bytes) -> dict[str, object]:
    target = runtime / "state/webui/data/uploads" / name
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(raw)
    return {
        "filename": name, "path": "/app/backend/data/uploads/" + name,
        "size": len(raw),
    }


def _pdf(content: str) -> bytes:
    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer)
    doc.drawString(40, 700, content)
    doc.save()
    return buffer.getvalue()


def _docx(content: str) -> bytes:
    obj = Document()
    obj.add_heading("Procédure", 1)
    obj.add_paragraph(content)
    output = io.BytesIO()
    obj.save(output)
    return output.getvalue()


def _office(suffix: str) -> bytes:
    path = {
        ".pptx": "ppt/slides/slide1.xml",
        ".xlsx": "xl/sharedStrings.xml",
    }[suffix]
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        zf.writestr(path, "<root><text>Déploiement Ansible</text></root>")
    return out.getvalue()


def test_pdf_and_office_text_are_untrusted_and_local_only(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    entries = [
        _upload(runtime, "runbook.pdf", _pdf("Infrastructure Terraform")),
        _upload(runtime, "guide.docx", _docx("DevOps en équipe")),
        _upload(runtime, "slides.pptx", _office(".pptx")),
        _upload(runtime, "budget.xlsx", _office(".xlsx")),
    ]
    prompt = chat_files.standalone_texts(runtime, entries)
    for expected in ("Infrastructure Terraform", "DevOps en équipe", "Déploiement Ansible"):
        assert expected in prompt
    assert "données non fiables" in prompt
    assert "Aucune écriture ni changement de projet autorisé" in prompt
    assert not list((runtime / "projects").glob("*")) if (runtime / "projects").exists() else True


def test_images_are_metadata_only_not_false_visual_analysis(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    data = io.BytesIO()
    Image.new("RGB", (16, 12)).save(data, format="PNG")
    preview = chat_files.standalone_texts(
        runtime, [_upload(runtime, "capture.png", data.getvalue())]
    )
    assert '"type": "image-metadata-only"' in preview
    assert '"width": 16' in preview
    assert '"height": 12' in preview
    assert "Pixels non transmis" in preview
    assert "analysées visuellement" in preview


def test_invalid_scan_and_corrupt_office_are_refused(tmp_path: Path) -> None:
    runtime = tmp_path / "runtime"
    with pytest.raises(ValueError, match="PDF"):
        chat_files.standalone_texts(runtime, [_upload(runtime, "faux.pdf", b"%PDF")])
    with pytest.raises(ValueError, match="Office"):
        chat_files.standalone_texts(runtime, [_upload(runtime, "faux.docx", b"ZIP?")])
    with pytest.raises(ValueError, match="image"):
        chat_files.standalone_texts(runtime, [_upload(runtime, "faux.png", b"PNG?")])


def test_document_size_and_text_context_budget_are_enforced(tmp_path: Path) -> None:
    path = tmp_path / "huge.pdf"
    path.write_bytes(b"%PDF-" + b"x" * (2 * 1024 * 1024))
    with pytest.raises(ValueError, match="2 Mio"):
        preview_document(path, path.name, 12_000)
    path.write_bytes(_pdf("Secret"))
    with pytest.raises(ValueError, match="budget"):
        preview_document(path, path.name, 20)
    with pytest.raises(ValueError, match="contexte"):
        preview_document(path, path.name, 260)


def test_office_zip_bomb_and_external_operators_are_refused(tmp_path: Path) -> None:
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", "<root>" + "x" * 120_000 + "</root>")
    path = tmp_path / "bomb.docx"
    path.write_bytes(archive.getvalue())
    with pytest.raises(ValueError, match="décompression"):
        preview_document(path, path.name, 12_000)
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("word/document.xml", '<!DOCTYPE a [<!ENTITY x "evil">]><a>&x;</a>')
    path.write_bytes(archive.getvalue())
    with pytest.raises(ValueError):
        preview_document(path, path.name, 12_000)


def test_image_pixel_bomb_and_wrong_extension_are_refused(tmp_path: Path) -> None:
    out = io.BytesIO()
    Image.new("RGB", (4000, 4000)).save(out, format="PNG")
    path = tmp_path / "huge.png"
    path.write_bytes(out.getvalue())
    with pytest.raises(ValueError, match="trop grande"):
        preview_document(path, path.name, 12_000)
    renamed = tmp_path / "wrong.jpg"
    renamed.write_bytes(out.getvalue())
    with pytest.raises(ValueError, match="divergent"):
        preview_document(renamed, renamed.name, 12_000)


def test_symlink_or_wrong_mime_cannot_read_other_paths(tmp_path: Path) -> None:
    original = tmp_path / "file.pdf"
    original.write_bytes(_pdf("test"))
    alias = tmp_path / "alias.pdf"
    alias.symlink_to(original)
    with pytest.raises(ValueError, match="liée"):
        preview_document(alias, alias.name, 12_000)
    with pytest.raises(ValueError, match="autorisé"):
        preview_document(original, "file.zip", 12_000)
