"""CPU-only, deterministic document exports from bounded Markdown, without resource fetches."""

from __future__ import annotations

import io
import textwrap
import zipfile
from datetime import datetime
from html import escape
from pathlib import Path
from typing import Any

from markdown_it import MarkdownIt

EXPORT_PREFIX = "@document-export:"
EXPORTS = {".pdf", ".docx", ".txt"}


def blocks(markdown: str) -> list[tuple[str, Any]]:
    if not isinstance(markdown, str) or not 0 < len(markdown.encode()) <= 60000:
        raise ValueError("source Markdown limitée à 60000 octets")
    tokens = MarkdownIt("commonmark", {"html": False}).enable("table").parse(markdown)
    result: list[tuple[str, Any]] = []
    row: list[str] = []
    table: list[list[str]] = []
    heading, listing, in_table = "paragraph", 0, False
    for token in tokens:
        if token.type == "heading_open":
            heading = token.tag
        elif token.type == "heading_close":
            heading = "paragraph"
        elif token.type in {"bullet_list_open", "ordered_list_open"}:
            listing += 1
        elif token.type in {"bullet_list_close", "ordered_list_close"}:
            listing -= 1
        elif token.type == "table_open":
            in_table, table = True, []
        elif token.type == "tr_open":
            row = []
        elif token.type == "tr_close":
            table.append(row)
        elif token.type == "table_close":
            in_table = False
            if len(table) > 100 or any(len(r) > 6 for r in table):
                raise ValueError("table limitée à 100 lignes et 6 colonnes")
            if any(len(c) > 3000 for r in table for c in r):
                raise ValueError("cellule de table limitée à 3000 caractères")
            result.append(("table", table))
        elif token.type == "inline":
            fragments: list[str] = []
            destinations: list[str] = []
            for child in token.children or []:
                if child.type == "link_open":
                    destinations.append(str(child.attrGet("href") or ""))
                elif child.type == "link_close" and destinations:
                    fragments.append(" (" + destinations.pop() + ")")
                elif child.type in {"text", "code_inline", "image"}:
                    fragments.append(child.content)
                elif child.type in {"softbreak", "hardbreak"}:
                    fragments.append("\n")
            text = "".join(fragments)
            if in_table:
                row.append(text)
            else:
                result.append((heading if not listing else "list", text))
        elif token.type in {"fence", "code_block"}:
            result.append(("code", token.content.rstrip()))
    if len(result) > 500:
        raise ValueError("500 blocs maximum par export; découper le document")
    return result


def export(markdown: str, kind: str) -> bytes:
    parts = blocks(markdown)
    if kind == "txt":
        return (
            "\n\n".join(
                "\n".join(" | ".join(row) for row in value) if block == "table" else value
                for block, value in parts
            )
            + "\n"
        ).encode()
    try:
        if kind == "docx":
            return _docx(parts)
        if kind == "pdf":
            return _pdf(parts)
    except Exception as exc:
        raise ValueError("export hors limites ou moteur indisponible") from exc
    raise ValueError("export pdf/docx/txt requis")


def output_bytes(relative: str, files: dict[str, str]) -> bytes:
    content = files[relative]
    suffix = Path(relative).suffix.lower()
    if content.startswith(EXPORT_PREFIX):
        source = content[len(EXPORT_PREFIX) :]
        if (
            suffix not in EXPORTS
            or source not in files
            or Path(source).suffix != ".md"
            or files[source].startswith(EXPORT_PREFIX)
            or Path(source).parent != Path(relative).parent
            or Path(source).stem != Path(relative).stem
        ):
            raise ValueError("export: source .md attendue dans la même tâche et le même dossier")
        return export(files[source], suffix[1:])
    if suffix in {".pdf", ".docx"}:
        raise ValueError(
            "PDF/DOCX: fournir une source .md et une référence d’export, pas du texte renommé"
        )
    return content.encode()


def _docx(parts: list[tuple[str, Any]]) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    doc.core_properties.created = doc.core_properties.modified = datetime(2000, 1, 1)
    doc.core_properties.author = "Atelier IA local"
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "DejaVu Sans", Pt(10.5)
    for kind, value in parts:
        if kind.startswith("h") and kind[1:].isdigit():
            doc.add_heading(value, min(int(kind[1:]), 6))
        elif kind == "table":
            table = doc.add_table(rows=0, cols=max(len(r) for r in value))
            table.style = "Light Shading Accent 1"
            for row in value:
                cells = table.add_row().cells
                for i, cell in enumerate(row):
                    cells[i].text = cell
        elif kind == "code":
            p = doc.add_paragraph(value)
            for run in p.runs:
                run.font.name, run.font.size = "DejaVu Sans Mono", Pt(8)
        else:
            doc.add_paragraph(value, style="List Bullet" if kind == "list" else None)
    buffer = io.BytesIO()
    doc.save(buffer)
    # ZIP metadata must be stable: reviewed bytes must equal the collected export.
    out = io.BytesIO()
    with (
        zipfile.ZipFile(buffer) as original,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for name in sorted(original.namelist()):
            entry = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(entry, original.read(name))
    return out.getvalue()


def _pdf(parts: list[tuple[str, Any]]) -> bytes:
    import reportlab  # type: ignore[import-untyped]
    from reportlab.lib import colors  # type: ignore[import-untyped]
    from reportlab.lib.enums import TA_LEFT  # type: ignore[import-untyped]
    from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
    from reportlab.lib.styles import ParagraphStyle  # type: ignore[import-untyped]
    from reportlab.lib.units import mm  # type: ignore[import-untyped]
    from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
    from reportlab.pdfbase.ttfonts import TTFont  # type: ignore[import-untyped]
    from reportlab.pdfgen.canvas import Canvas  # type: ignore[import-untyped]
    from reportlab.platypus import (  # type: ignore[import-untyped]
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font_dir = Path(reportlab.__file__).parent / "fonts"
    # Bundled fonts make outputs independent of desktop packages and reproducible in CI.
    for name, file in (("Atelier", "Vera.ttf"), ("AtelierBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(font_dir / file)))
    styles = {
        "body": ParagraphStyle(
            "body",
            fontName="Atelier",
            fontSize=10,
            leading=15,
            spaceAfter=8,
            alignment=TA_LEFT,
            splitLongWords=True,
        ),
        "heading": ParagraphStyle(
            "heading",
            fontName="AtelierBold",
            fontSize=17,
            leading=22,
            spaceAfter=10,
            spaceBefore=12,
            textColor=colors.HexColor("#126b74"),
        ),
        "code": ParagraphStyle(
            "code",
            fontName="Atelier",
            fontSize=8,
            leading=11,
            spaceAfter=8,
            backColor=colors.HexColor("#eef3f6"),
        ),
    }
    story = []
    for kind, value in parts:
        if kind == "table":
            columns = max(len(r) for r in value)
            table = Table(
                [[Paragraph(escape(c), styles["body"]) for c in row] for row in value],
                colWidths=[170 * mm / columns] * columns,
                repeatRows=1,
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf5f3")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
                    ]
                )
            )
            story.extend([table, Spacer(1, 8)])
        elif kind == "code":
            # Wrap for print only; the Markdown source keeps exact executable lines.
            for line in value.splitlines() or [""]:
                for wrapped in textwrap.wrap(
                    line, 88, replace_whitespace=False, drop_whitespace=False
                ) or [""]:
                    story.append(
                        Paragraph(escape(wrapped).replace(" ", "&#160;") or "&#160;", styles["code"])
                    )
        else:
            style = styles["heading"] if kind.startswith("h") else styles["body"]
            story.append(
                Paragraph(
                    ("• " if kind == "list" else "") + escape(value).replace("\n", "<br/>"), style
                )
            )
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title="Atelier IA local",
        author="Atelier IA local",
    )

    def footer(canvas: Any, document: Any) -> None:
        if document.page > 80:
            raise ValueError("80 pages maximum; découper le document")
        canvas.setFont("Atelier", 8)
        canvas.drawString(20 * mm, 12 * mm, "Atelier IA local · Infrastructure & OPS")
        canvas.drawRightString(190 * mm, 12 * mm, str(document.page))

    def stable_canvas(*args: Any, **kwargs: Any) -> Any:
        kwargs["invariant"] = 1
        return Canvas(*args, **kwargs)

    doc.build(story, onFirstPage=footer, onLaterPages=footer, canvasmaker=stable_canvas)
    return buffer.getvalue()
